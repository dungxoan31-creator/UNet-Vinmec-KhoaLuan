import os
import sys
import hashlib
import docx

# Ensure utf-8 output
sys.stdout.reconfigure(encoding='utf-8')

reports_dir = 'docs/reports'
print("--- DOCX INSPECTION IN docs/reports ---")
for f in sorted(os.listdir(reports_dir)):
    p = os.path.join(reports_dir, f)
    if os.path.isfile(p):
        sz = os.path.getsize(p)
        mtime = os.path.getmtime(p)
        if f.startswith('~$'):
            print(f"[TEMP LOCK FILE] {f} ({sz} bytes)")
            continue
        if f.endswith('.docx'):
            doc = docx.Document(p)
            paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            tables = len(doc.tables)
            total_words = sum(len(p.split()) for p in paras)
            # check images inside docx
            rel_imgs = [rel.target_ref for rel in doc.part.rels.values() if "image" in rel.reltype]
            print(f"\nFile: {f}")
            print(f"  Size: {sz:,} bytes | Tables: {tables} | Paragraphs: {len(paras)} | Words: {total_words} | Embedded Images: {len(rel_imgs)}")
            print(f"  Title / First para: {paras[0] if paras else '[EMPTY]'}")
            if len(paras) > 1:
                print(f"  Second para: {paras[1][:120]}...")
