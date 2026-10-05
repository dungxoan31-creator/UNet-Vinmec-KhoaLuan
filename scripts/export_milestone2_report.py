"""Export the milestone 2 Markdown report to a Word document."""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/reports/Bao_Cao_Tien_Do_Moc_2_NguyenHuuDung_2026-10-03.md"
OUTPUT = ROOT / "docs/reports/BaoCaoTienDoMoc2_NguyenHuuDung_11235559_final.docx"
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))")
IMAGE = re.compile(r"!\[([^]]+)\]\(([^)]+)\)")


def add_hyperlink(paragraph, label, target):
    relation = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "17518A")
    properties.append(color)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, content):
    for piece in INLINE.split(content):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            paragraph.add_run(piece[1:-1])
        elif piece.startswith("[") and "](" in piece:
            label, target = piece[1:-1].split("](", 1)
            add_hyperlink(paragraph, label, target)
        else:
            paragraph.add_run(piece.replace("  \n", "\n"))


def table_cells(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def style_table_cell(cell, row_index):
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "5")
        border.set(qn("w:color"), "B7C0C8")
        borders.append(border)
    properties.append(borders)
    if row_index == 0 or row_index % 2 == 0:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E9EDF1" if row_index == 0 else "F7F9FA")
        properties.append(shading)


def add_table(document, lines):
    rows = [table_cells(line) for line in lines]
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            style_table_cell(cell, row_index)
            paragraph = cell.paragraphs[0]
            if row_index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8.5 if len(rows[0]) >= 5 else 9.5)
                if row_index == 0:
                    run.bold = True
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
    return table


def export():
    document = Document()
    document.core_properties.title = "Báo cáo tiến độ Mốc 2 – Nguyễn Hữu Dũng"
    document.core_properties.author = "Nguyễn Hữu Dũng"
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(3), Cm(2)
    section.top_margin = section.bottom_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.35
    for style_name, size in (("Heading 1", 14), ("Heading 2", 12.5)):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Trang ")
    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(10)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    paragraph_lines = []

    def flush_paragraph():
        if not paragraph_lines:
            return
        content = "\n".join(paragraph_lines).replace("  \n", "\n")
        paragraph = document.add_paragraph()
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
            if content.startswith(("Trường Công nghệ", "***", "**MỐC 2"))
            else WD_ALIGN_PARAGRAPH.JUSTIFY
        )
        add_inline(paragraph, content)
        paragraph_lines.clear()

    while index < len(lines):
        line = lines[index].strip()
        if not line:
            flush_paragraph()
            index += 1
            continue
        if line == "---":
            flush_paragraph()
            index += 1
            continue
        image = IMAGE.fullmatch(line)
        if image:
            flush_paragraph()
            image_path = (SOURCE.parent / image.group(2)).resolve()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image_path), width=Inches(6.5))
            caption = document.add_paragraph(image.group(1))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            index += 1
            continue
        if line.startswith("|"):
            flush_paragraph()
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            add_table(document, block)
            continue
        if line.startswith("# ") or line.startswith("## ") or line.startswith("### "):
            flush_paragraph()
            title = line.lstrip("# ")
            style = "Heading 1" if line.startswith("## ") else "Heading 2" if line.startswith("### ") else None
            paragraph = document.add_paragraph(style=style)
            paragraph.paragraph_format.keep_with_next = True
            if line.startswith("# "):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                size = 16 if title.startswith("BÁO CÁO") else 13
                paragraph.paragraph_format.space_before = Pt(8)
            elif line.startswith("## "):
                size = 14
            else:
                size = 12.5
            run = paragraph.add_run(title)
            run.bold = True
            run.font.size = Pt(size)
            index += 1
            continue
        if line.startswith("**Bảng "):
            flush_paragraph()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, line)
            index += 1
            continue
        if line.startswith("- "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
            index += 1
            continue
        paragraph_lines.append(lines[index])
        index += 1
    flush_paragraph()
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(export())
