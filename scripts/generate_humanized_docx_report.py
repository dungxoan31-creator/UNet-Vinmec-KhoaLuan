import os
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in dxa (1 pt = 20 dxa)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F2F4F7"):
    """Set cell background color."""
    shading_xml = f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table):
    """Set elegant, academic borders for a table."""
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="8" w:space="0" w:color="4A5568"/>
        <w:bottom w:val="single" w:sz="8" w:space="0" w:color="4A5568"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E0"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_footer_page_number(run):
    """Add standard Word page number field to a run."""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def create_report():
    doc = docx.Document()
    
    # -------------------------------------------------------------
    # 1. Page Setup: A4, Margins: Left 3.0cm, Right 2.0cm, Top 2.0cm, Bottom 2.0cm
    # -------------------------------------------------------------
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    
    # Footer: Page number centered
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.first_line_indent = Cm(0)
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    f_run = footer_p.add_run()
    f_run.font.name = 'Times New Roman'
    f_run.font.size = Pt(10)
    f_run.font.color.rgb = RGBColor(100, 100, 100)
    add_footer_page_number(f_run)

    # -------------------------------------------------------------
    # Helper functions for adding elements with strict formatting
    # -------------------------------------------------------------
    def add_p(text="", style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, 
              first_indent=Cm(1.27), space_before=Pt(0), space_after=Pt(6), 
              line_spacing=1.5, bold=False, italic=False, font_size=Pt(13), 
              color=RGBColor(0, 0, 0)):
        p = doc.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.first_line_indent = first_indent
        pf.space_before = space_before
        pf.space_after = space_after
        pf.line_spacing = line_spacing
        
        if text:
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = font_size
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(12)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.3
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.3
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(6)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.3
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_table_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(8)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.2
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_table_source(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(2)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.2
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10.5)
        run.font.italic = True
        run.font.color.rgb = RGBColor(80, 80, 80)
        return p

    def add_figure_caption(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(3)
        pf.space_after = Pt(8)
        pf.line_spacing = 1.2
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        run.font.italic = True
        run.font.color.rgb = RGBColor(40, 40, 40)
        return p

    # -------------------------------------------------------------
    # Document Header & Metadata Block
    # -------------------------------------------------------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.first_line_indent = Cm(0)
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(2)
    run_inst = p_inst.add_run("TRƯỜNG ĐẠI HỌC KINH TẾ QUỐC DÂN\nTRƯỜNG CÔNG NGHỆ VÀ KINH TẾ SỐ")
    run_inst.font.name = 'Times New Roman'
    run_inst.font.size = Pt(12)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(0, 0, 0)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.first_line_indent = Cm(0)
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("BÁO CÁO TIẾN ĐỘ THỰC HIỆN KHÓA LUẬN TỐT NGHIỆP\nCỘT MỐC 1: DỮ LIỆU, TIỀN XỬ LÝ VÀ MÔ HÌNH PHÂN ĐOẠN CƠ SỞ")
    run_title.font.name = 'Times New Roman'
    run_title.font.size = Pt(15)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.first_line_indent = Cm(0)
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("Thời gian thực hiện: Từ ngày 06/09/2026 đến ngày 20/09/2026")
    run_sub.font.name = 'Times New Roman'
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(80, 80, 80)

    # Information block
    info_items = [
        ("Tên đề tài: ", "Xây dựng hệ thống hỗ trợ phân đoạn tổn thương trên ảnh siêu âm buồng trứng ứng dụng Deep Learning theo mô hình Human-in-the-Loop."),
        ("Sinh viên thực hiện: ", "Nguyễn Hữu Dũng (Mã sinh viên: 11235559)."),
        ("Lớp chuyên ngành: ", "Hệ thống Thông tin Quản lý 65A (HTTTQL 65A)."),
        ("Giảng viên hướng dẫn: ", "ThS. Trần Thanh Hải."),
        ("Ngày lập báo cáo: ", "17/09/2026.")
    ]
    for label, val in info_items:
        p_info = doc.add_paragraph()
        p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_info.paragraph_format.first_line_indent = Cm(0)
        p_info.paragraph_format.space_before = Pt(0)
        p_info.paragraph_format.space_after = Pt(3)
        p_info.paragraph_format.line_spacing = 1.3
        r_lbl = p_info.add_run(label)
        r_lbl.font.name = 'Times New Roman'
        r_lbl.font.size = Pt(12)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(0, 0, 0)
        r_val = p_info.add_run(val)
        r_val.font.name = 'Times New Roman'
        r_val.font.size = Pt(12)
        r_val.font.color.rgb = RGBColor(0, 0, 0)

    # Separator spacing
    add_p("", space_after=Pt(6))

    # Opening letter
    add_p("Kính gửi: ThS. Trần Thanh Hải,", first_indent=Cm(0), bold=True)
    add_p("Theo đề cương nghiên cứu đã được Thầy thông qua với phạm vi trọng tâm: Ảnh siêu âm -> Tiền xử lý -> Phân đoạn tổn thương -> Lớp phủ mặt nạ -> Bác sĩ đánh giá và chỉnh sửa -> Đánh giá kết quả.")
    add_p("Em xin gửi Thầy báo cáo chi tiết về kết quả thực hiện Cột mốc 1 trong giai đoạn từ ngày 06/09/2026 đến ngày 20/09/2026. Nội dung báo cáo bao gồm việc hoàn thành rà soát tập dữ liệu lâm sàng thu nhận tại Bệnh viện Đa khoa Quốc tế Vinmec Times City, thiết lập chuỗi tiền xử lý ảnh y tế, xây dựng bộ nạp dữ liệu và kịch bản kiểm thử tự động, huấn luyện mô hình Standard U-Net cơ sở trên phần cứng thực tế và đánh giá định lượng độc lập trên tập kiểm thử 46 ca.")

    # -------------------------------------------------------------
    # SECTION I
    # -------------------------------------------------------------
    add_heading_1("I. CÁC HẠNG MỤC CÔNG VIỆC ĐÃ HOÀN THÀNH")
    add_p("Trong Cột mốc 1, toàn bộ sáu nhiệm vụ kỹ thuật đã được triển khai tuần tự theo kế hoạch và hoàn tất việc kiểm tra thực nghiệm trên hệ thống:")

    add_table_title("Bảng 1. Danh mục các nhiệm vụ kỹ thuật đã hoàn thành trong Cột mốc 1")
    t1_data = [
        ["STT", "Nhiệm vụ kỹ thuật", "Nội dung triển khai", "Tệp mã nguồn và dữ liệu liên quan", "Kết quả"],
        ["1", "Khảo sát và đóng băng dữ liệu Ground Truth", 
         "Rà soát kho 1.387 ảnh thô từ BVĐKQT Vinmec Times City; chốt danh mục 307 ảnh Ground Truth từ 185 bệnh nhân; xác định 35 ca mask rỗng (11,4%) phục vụ kiểm chứng âm tính.",
         "dataset/vinmec_ovarian/vinmec_dataset_manifest.json\nmetadata/kltn_ground_truth_307.csv",
         "Đã đóng băng"],
        ["2", "Phân chia dữ liệu theo cấp bệnh nhân (Patient-Level Split)",
         "Cài đặt thuật toán phân tầng theo mã bệnh nhân; kiểm định loại trừ 100% việc trùng lặp bệnh nhân giữa ba tập; tỷ lệ phân chia Train (215 ảnh), Val (46 ảnh) và Test (46 ảnh).",
         "scripts/prepare_splits.py\nai_training/splits/train.csv\nai_training/splits/val.csv\nai_training/splits/test.csv",
         "Không rò rỉ dữ liệu"],
        ["3", "Xây dựng pipeline tiền xử lý chuẩn y tế (Medical Preprocessor)",
         "Cắt vùng quạt quét siêu âm (ROI Fan-beam); kỹ thuật Letterbox kích thước 512x512 giữ nguyên tỷ lệ hình thái; bảo toàn nhãn nhị phân {0, 1} bằng phép nội suy lân cận gần nhất; lập trình hàm khôi phục tọa độ gốc inverse_letterbox_mask.",
         "backend/services/preprocessor.py\ntests/test_milestone_1_audit.py",
         "Đạt kiểm thử"],
        ["4", "Xây dựng PyTorch DataLoader và bộ kiểm thử tự động",
         "Xây dựng lớp nạp dữ liệu kèm tăng cường hình ảnh an toàn trên tập huấn luyện; cài đặt các hàm đo lường chỉ số lâm sàng; xây dựng bộ kiểm thử 16 tiêu chí tự động.",
         "ai_training/dataset_loader.py\nai_training/metrics_clinical.py\ntests/test_milestone_1_audit.py",
         "16/16 bài kiểm thử đạt (7,31s)"],
        ["5", "Huấn luyện mô hình Standard U-Net cơ sở",
         "Huấn luyện kiến trúc U-Net 7,76 triệu tham số trên GPU NVIDIA RTX 3050; sử dụng Combo Loss kết hợp BCE và Soft Dice; áp dụng bộ tối ưu AdamW cùng lịch học Cosine Annealing; lưu checkpoint tại epoch 12.",
         "backend/models/unet.py\nai_training/train_baseline_unet.py\ncheckpoints/baseline_unet_best.pth",
         "Validation Dice: 0,5830"],
        ["6", "Đánh giá định lượng trên tập kiểm thử độc lập",
         "Đánh giá độc lập trên 46 ảnh của tập kiểm thử; tính toán các chỉ số Dice, IoU, Recall, Precision, Specificity; kết xuất đồ thị huấn luyện và năm bảng ảnh trực quan đối chứng.",
         "evaluation/evaluate_baseline_testset.py\nevaluation/baseline_test_metrics.json\nevaluation/baseline_visualizations/",
         "Foreground Dice: 0,5719\nIoU: 0,4405"]
    ]
    
    table1 = doc.add_table(rows=len(t1_data), cols=5)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table1)
    col_widths1 = [Cm(1.0), Cm(3.2), Cm(5.8), Cm(3.8), Cm(2.2)]
    
    for r_idx, row in enumerate(table1.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths1[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p_c = cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c_idx in [0, 4] or r_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
            p_c.paragraph_format.first_line_indent = Cm(0)
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(0)
            p_c.paragraph_format.line_spacing = 1.15
            run_c = p_c.add_run(t1_data[r_idx][c_idx])
            run_c.font.name = 'Times New Roman'
            run_c.font.size = Pt(10.5) if r_idx > 0 else Pt(11)
            run_c.font.bold = (r_idx == 0 or c_idx == 0)
            if r_idx == 0:
                set_cell_shading(cell, "EDF2F7")
                run_c.font.color.rgb = RGBColor(26, 32, 44)
    
    add_table_source("Nguồn: Tổng hợp từ kế hoạch thực hiện đề cương khóa luận tốt nghiệp.")

    # -------------------------------------------------------------
    # SECTION II
    # -------------------------------------------------------------
    add_heading_1("II. DỮ LIỆU NGHIÊN CỨU VÀ QUY TRÌNH TIỀN XỬ LÝ CHUẨN Y KHOA")
    add_p("Trong bài toán chẩn đoán hình ảnh y tế, việc phân chia dữ liệu ngẫu nhiên theo từng ảnh chụp rất dễ dẫn đến hiện tượng rò rỉ dữ liệu (data leakage) khi ảnh của cùng một bệnh nhân xuất hiện ở cả tập huấn luyện và tập kiểm thử. Hiện tượng này làm cho mô hình ghi nhớ đặc điểm mô học cục bộ của bệnh nhân thay vì học các đặc trưng bệnh lý phổ quát, dẫn đến việc đánh giá độ chính xác bị thổi phồng so với thực tế lâm sàng.")
    add_p("Để giải quyết triệt để rủi ro trên, thuật toán phân chia trong đề tài được thiết lập theo cấp bệnh nhân (Patient-Level Split). Toàn bộ 307 ảnh Ground Truth từ 185 bệnh nhân được phân chia thành ba tập dữ liệu độc lập hoàn toàn với tỷ lệ 70% huấn luyện, 15% thẩm định và 15% kiểm thử.")

    add_table_title("Bảng 2. Cơ cấu phân chia tập dữ liệu nghiên cứu theo cấp bệnh nhân")
    t2_data = [
        ["Tập dữ liệu", "Số lượng ảnh", "Tỷ lệ ảnh (%)", "Số bệnh nhân độc lập", "Số ca mask rỗng", "Mục đích sử dụng"],
        ["Tập huấn luyện (Train set)", "215", "70,03%", "130", "25", "Cập nhật trọng số của mô hình U-Net"],
        ["Tập thẩm định (Val set)", "46", "14,98%", "27", "5", "Tinh chỉnh siêu tham số và lưu checkpoint"],
        ["Tập kiểm thử (Test set)", "46", "14,98%", "28", "5", "Đánh giá độc lập năng lực mô hình"],
        ["Tổng cộng (Ground Truth)", "307", "100,0%", "185", "35 (11,4%)", "Tập dữ liệu được niêm phong nghiên cứu"]
    ]
    table2 = doc.add_table(rows=len(t2_data), cols=6)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table2)
    col_widths2 = [Cm(3.8), Cm(2.0), Cm(2.0), Cm(2.4), Cm(2.4), Cm(3.4)]
    
    for r_idx, row in enumerate(table2.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths2[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p_c = cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c_idx in [1, 2, 3, 4] or r_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
            p_c.paragraph_format.first_line_indent = Cm(0)
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(0)
            p_c.paragraph_format.line_spacing = 1.15
            run_c = p_c.add_run(t2_data[r_idx][c_idx])
            run_c.font.name = 'Times New Roman'
            run_c.font.size = Pt(10.5) if r_idx > 0 else Pt(11)
            run_c.font.bold = (r_idx == 0 or r_idx == len(t2_data)-1)
            if r_idx == 0:
                set_cell_shading(cell, "EDF2F7")
            elif r_idx == len(t2_data)-1:
                set_cell_shading(cell, "F7FAFC")
    
    add_table_source("Nguồn: Trích xuất từ tệp de_cuong_split_manifest.json và vinmec_dataset_manifest.json.")
    add_p("Kiểm toán mã nguồn xác nhận tập giao bệnh nhân giữa ba tập là rỗng: Train ∩ Val = ∅, Train ∩ Test = ∅, Val ∩ Test = ∅. Tỷ lệ ca mask rỗng (11,4%) được phân bổ đồng đều ở cả ba tập nhằm kiểm chứng khả năng phát hiện âm tính và tránh báo động giả trên các ca không có tổn thương.")

    add_heading_2("Quy trình tiền xử lý ảnh năm bước")
    add_p("Ảnh siêu âm buồng trứng thường đi kèm nhiều yếu tố gây nhiễu như vùng quét hình quạt, thông tin văn bản hành chính của máy, thước đo độ sâu và dải tương phản không đồng đều. Chuỗi tiền xử lý trong hệ thống được thiết kế thành năm bước nối tiếp:")
    add_p("1. Đọc ảnh thang độ xám và trích xuất mặt nạ tổn thương dạng nhị phân;", first_indent=Cm(1.27))
    add_p("2. Tự động xác định và cắt vùng quạt quét sóng âm (ROI Fan-beam) để loại bỏ toàn bộ khung viền đen và thông số văn bản xung quanh;", first_indent=Cm(1.27))
    add_p("3. Áp dụng kỹ thuật Letterbox đưa ảnh về kích thước chuẩn 512x512 pixel, bổ sung đệm viền đen đối xứng nhằm giữ nguyên tỷ lệ khung hình tự nhiên của khối u. Phép nội suy song tuyến tính (Bilinear) được áp dụng cho ảnh và phép nội suy lân cận gần nhất (Nearest-Neighbor) được áp dụng cho mặt nạ nhằm bảo toàn nhãn nhị phân {0, 1};", first_indent=Cm(1.27))
    add_p("4. Áp dụng thuật toán cân bằng tương phản thích ứng cục bộ (CLAHE) với ngưỡng giới hạn clip_limit = 2.0 để tăng cường độ tương phản giữa ranh giới u nang và mô đệm buồng trứng;", first_indent=Cm(1.27))
    add_p("5. Cài đặt hàm khôi phục kích thước ngược (inverse_letterbox_mask) nhằm ánh xạ chính xác kết quả phân đoạn 512x512 pixel trở lại hệ tọa độ không gian ban đầu của ảnh siêu âm.", first_indent=Cm(1.27))

    # Insert Image 1
    img1_path = "evaluation/baseline_visualizations/preprocessing_pipeline_demonstration.png"
    if os.path.exists(img1_path):
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.paragraph_format.first_line_indent = Cm(0)
        p_img1.paragraph_format.space_before = Pt(6)
        p_img1.paragraph_format.space_after = Pt(2)
        p_img1.add_run().add_picture(img1_path, width=Inches(6.0))
        add_figure_caption("Hình 1. Quy trình tiền xử lý ảnh siêu âm buồng trứng năm bước từ ảnh thô đến khôi phục tọa độ gốc")

    # -------------------------------------------------------------
    # SECTION III
    # -------------------------------------------------------------
    add_heading_1("III. THÔNG SỐ HUẤN LUYỆN VÀ NHẬT KÝ HỘI TỤ MÔ HÌNH")
    add_p("Mô hình phân đoạn cơ sở được xây dựng dựa trên kiến trúc Standard U-Net nguyên bản (Ronneberger và cộng sự, 2015). Mạng gồm 4 tầng mã hóa (Encoder) trích xuất đặc trưng với số kênh tăng dần (32 -> 64 -> 128 -> 256), 1 tầng nghẽn cổ chai (Bottleneck, 512 kênh) và 4 tầng giải mã (Decoder) khôi phục độ phân giải không gian thông qua các đường truyền tắt (skip connections).")

    add_table_title("Bảng 3. Thông số kỹ thuật của mô hình Standard U-Net cơ sở")
    t3_data = [
        ["Tham số kỹ thuật", "Giá trị thiết lập", "Ghi chú kỹ thuật và phần cứng"],
        ["Kiến trúc mô hình", "Standard U-Net (Ronneberger et al., 2015)", "4 tầng Encoder, 1 Bottleneck, 4 tầng Decoder, liên kết tắt"],
        ["Tổng số tham số", "7.762.465 tham số (~7,76 triệu)", "Base filters: 32 -> 64 -> 128 -> 256 -> 512"],
        ["Kích thước đầu vào", "(B, 1, 512, 512) kiểu float32", "Định dạng chuẩn hóa qua quy trình Letterbox"],
        ["Hàm mất mát (Loss)", "Combo Loss (0,5 BCE + 0,5 Soft Dice)", "Cân bằng giữa tối ưu phân loại pixel và độ trùng lặp hình học"],
        ["Bộ tối ưu hóa", "AdamW (lr = 1e-3, weight_decay = 1e-4)", "Tối ưu hóa có điều chuẩn L2 ổn định"],
        ["Lịch điều chỉnh tốc độ học", "Cosine Annealing (T_max = 12, eta_min = 1e-6)", "Giảm dần tốc độ học theo chu kỳ cosine"],
        ["Kỹ thuật tính toán", "Automatic Mixed Precision (AMP FP16)", "Giảm chiếm dụng bộ nhớ, tăng tốc tính toán trên GPU RTX 3050"],
        ["Kích thước lô (Batch size)", "batch_size = 4", "Mức tiêu hao bộ nhớ ~1,4 GB VRAM, tránh tràn RAM hệ thống"]
    ]
    table3 = doc.add_table(rows=len(t3_data), cols=3)
    table3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table3)
    col_widths3 = [Cm(4.0), Cm(4.8), Cm(7.2)]
    
    for r_idx, row in enumerate(table3.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths3[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p_c = cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c_idx == 1 and r_idx > 0) else WD_ALIGN_PARAGRAPH.LEFT
            p_c.paragraph_format.first_line_indent = Cm(0)
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(0)
            p_c.paragraph_format.line_spacing = 1.15
            run_c = p_c.add_run(t3_data[r_idx][c_idx])
            run_c.font.name = 'Times New Roman'
            run_c.font.size = Pt(10.5) if r_idx > 0 else Pt(11)
            run_c.font.bold = (r_idx == 0 or c_idx == 0)
            if r_idx == 0:
                set_cell_shading(cell, "EDF2F7")
    
    add_table_source("Nguồn: Cấu hình mã nguồn tại backend/models/unet.py và ai_training/train_baseline_unet.py.")

    add_p("Mô hình được huấn luyện trong 12 epoch trên máy tính xách tay cá nhân trang bị GPU NVIDIA GeForce RTX 3050 Laptop. Quá trình hội tụ diễn ra ổn định, hàm mất mát giảm dần qua các chu kỳ và hệ số Dice trên tập thẩm định liên tục được cải thiện, đạt giá trị cao nhất 0,5830 tại epoch 12.")

    add_table_title("Bảng 4. Nhật ký huấn luyện mô hình qua 12 epoch")
    t4_data = [
        ["Epoch", "Train Loss", "Train Dice", "Val Loss", "Val Dice", "Tốc độ học (LR)", "Ghi nhận Checkpoint"],
        ["Epoch 01", "0.6072", "0.3306", "1.4062", "0.3562", "0.000983", "Lưu checkpoint (Best Val Dice)"],
        ["Epoch 02", "0.5150", "0.4184", "0.5107", "0.4321", "0.000933", "Lưu checkpoint (Best Val Dice)"],
        ["Epoch 03", "0.4738", "0.4588", "0.5841", "0.4309", "0.000854", "-"],
        ["Epoch 04", "0.4313", "0.5083", "0.6683", "0.4427", "0.000750", "Lưu checkpoint (Best Val Dice)"],
        ["Epoch 05", "0.3986", "0.5365", "0.5040", "0.5239", "0.000630", "Lưu checkpoint (Best Val Dice)"],
        ["Epoch 06", "0.3845", "0.5467", "0.4341", "0.5053", "0.000501", "-"],
        ["Epoch 07", "0.3611", "0.5789", "0.4158", "0.5188", "0.000371", "-"],
        ["Epoch 08", "0.3468", "0.5902", "0.3815", "0.5615", "0.000251", "Lưu checkpoint (Best Val Dice)"],
        ["Epoch 09", "0.3356", "0.6068", "0.4039", "0.5556", "0.000147", "-"],
        ["Epoch 10", "0.3291", "0.6183", "0.4910", "0.5046", "0.000068", "-"],
        ["Epoch 11", "0.3153", "0.6385", "0.3728", "0.5751", "0.000018", "Lưu checkpoint (Best Val Dice)"],
        ["Epoch 12", "0.3204", "0.6259", "0.3762", "0.5830", "0.000001", "Lưu checkpoint tối ưu (0,5830)"]
    ]
    table4 = doc.add_table(rows=len(t4_data), cols=7)
    table4.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table4)
    col_widths4 = [Cm(2.0), Cm(2.2), Cm(2.2), Cm(2.2), Cm(2.2), Cm(2.4), Cm(2.8)]
    
    for r_idx, row in enumerate(table4.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths4[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            p_c = cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_c.paragraph_format.first_line_indent = Cm(0)
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(0)
            p_c.paragraph_format.line_spacing = 1.15
            run_c = p_c.add_run(t4_data[r_idx][c_idx])
            run_c.font.name = 'Times New Roman'
            run_c.font.size = Pt(10) if r_idx > 0 else Pt(10.5)
            run_c.font.bold = (r_idx == 0 or r_idx == len(t4_data)-1)
            if r_idx == 0:
                set_cell_shading(cell, "EDF2F7")
            elif r_idx == len(t4_data)-1:
                set_cell_shading(cell, "F7FAFC")
    
    add_table_source("Nguồn: Trích xuất từ nhật ký huấn luyện tại checkpoints/baseline_training_history.json.")

    # Insert Image 2
    img2_path = "evaluation/baseline_visualizations/training_convergence_curves.png"
    if os.path.exists(img2_path):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.first_line_indent = Cm(0)
        p_img2.paragraph_format.space_before = Pt(6)
        p_img2.paragraph_format.space_after = Pt(2)
        p_img2.add_run().add_picture(img2_path, width=Inches(5.8))
        add_figure_caption("Hình 2. Đồ thị hội tụ hàm mất mát Combo Loss và hệ số Dice trên tập huấn luyện và thẩm định qua 12 epoch")

    # -------------------------------------------------------------
    # SECTION IV
    # -------------------------------------------------------------
    add_heading_1("IV. KẾT QUẢ ĐÁNH GIÁ ĐỊNH LƯỢNG VÀ PHÂN TÍCH HÌNH ẢNH THỰC NGHIỆM")
    add_p("Mô hình Standard U-Net tốt nhất (lưu tại epoch 12) được đánh giá độc lập trên toàn bộ 46 ca của tập kiểm thử. Quá trình đánh giá áp dụng các chỉ số đo lường hình học và lâm sàng chuyên dụng:")

    add_table_title("Bảng 5. Kết quả đánh giá định lượng trên tập kiểm thử độc lập (46 ca)")
    t5_data = [
        ["Chỉ số đo lường", "Kết quả thực nghiệm", "Ý nghĩa chuyên môn và đánh giá y tế"],
        ["Foreground Dice (DSC)", "0,5719 ± 0,2482", "Mức độ trùng lặp không gian giữa vùng u dự đoán và Ground Truth do chuyên gia xác nhận."],
        ["Intersection over Union (IoU)", "0,4405 ± 0,2341", "Chỉ số Jaccard đo độ khít ranh giới vùng tổn thương buồng trứng."],
        ["Recall (Độ nhạy)", "0,6362 (63,6%)", "Tỷ lệ diện tích u thực tế được mô hình phát hiện, phản ánh khả năng hạn chế bỏ sót tổn thương."],
        ["Precision (Độ chuẩn xác)", "0,5864 (58,6%)", "Tỷ lệ diện tích dự đoán thực sự là tổn thương, phản ánh mức độ lan ra mô lành lân cận."],
        ["Specificity trên toàn bộ ảnh", "0,8977 (89,8%)", "Khả năng nhận diện chính xác các vùng mô lành xung quanh trên toàn khung hình."],
        ["Specificity trên ca mask rỗng", "0,8235 (82,4%)", "Khả năng nhận biết chính xác các ca không có u, hạn chế báo động giả."]
    ]
    table5 = doc.add_table(rows=len(t5_data), cols=3)
    table5.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table5)
    col_widths5 = [Cm(4.5), Cm(3.2), Cm(8.3)]
    
    for r_idx, row in enumerate(table5.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths5[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p_c = cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c_idx == 1 and r_idx > 0) else WD_ALIGN_PARAGRAPH.LEFT
            p_c.paragraph_format.first_line_indent = Cm(0)
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(0)
            p_c.paragraph_format.line_spacing = 1.15
            run_c = p_c.add_run(t5_data[r_idx][c_idx])
            run_c.font.name = 'Times New Roman'
            run_c.font.size = Pt(10.5) if r_idx > 0 else Pt(11)
            run_c.font.bold = (r_idx == 0 or c_idx == 0)
            if r_idx == 0:
                set_cell_shading(cell, "EDF2F7")
    
    add_table_source("Nguồn: Kết xuất từ kịch bản evaluation/evaluate_baseline_testset.py và tệp evaluation/baseline_test_metrics.json.")

    add_heading_2("Phân tích trực quan các trường hợp phân đoạn thực tế")
    add_p("Quan sát kết quả đối chứng trực quan trên 46 ca kiểm thử cho thấy sự phân hóa rõ nét giữa các nhóm hình thái bệnh lý:")
    add_p("1. Nhóm phân đoạn có độ chính xác cao (Dice > 0,75): Gồm các ca u nang bì hoặc nang đơn thùy có ranh giới âm học rõ ràng, cấu trúc cản âm phân biệt tốt với mô buồng trứng lành xung quanh. Mô hình Standard U-Net bám sát đường viền Ground Truth của bác sĩ.")

    # Insert Image 3
    img3_path = "evaluation/baseline_visualizations/best_matches.png"
    if os.path.exists(img3_path):
        p_img3 = doc.add_paragraph()
        p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img3.paragraph_format.first_line_indent = Cm(0)
        p_img3.paragraph_format.space_before = Pt(6)
        p_img3.paragraph_format.space_after = Pt(2)
        p_img3.add_run().add_picture(img3_path, width=Inches(5.0))
        add_figure_caption("Hình 3. Các trường hợp phân đoạn đạt độ chính xác cao (Dice > 0,75)")

    add_p("2. Nhóm phân đoạn có độ chính xác trung bình (Dice từ 0,60 đến 0,70): Mô hình định vị chính xác vị trí tâm và hình thái chung của khối u, tuy nhiên ranh giới phân đoạn ngoài rìa hơi lan nhẹ sang các cấu trúc lân cận có mức độ phản hồi âm thấp.")

    # Insert Image 4
    img4_path = "evaluation/baseline_visualizations/average_matches.png"
    if os.path.exists(img4_path):
        p_img4 = doc.add_paragraph()
        p_img4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img4.paragraph_format.first_line_indent = Cm(0)
        p_img4.paragraph_format.space_before = Pt(6)
        p_img4.paragraph_format.space_after = Pt(2)
        p_img4.add_run().add_picture(img4_path, width=Inches(5.0))
        add_figure_caption("Hình 4. Các trường hợp phân đoạn đạt độ chính xác trung bình (Dice từ 0,60 đến 0,70)")

    add_p("3. Nhóm ca bệnh khó và thách thức phân đoạn: Gồm các ca tổn thương buồng trứng bị che khuất bởi vệt bóng cản âm (acoustic shadowing) xuất phát từ thành phần tóc, bã đậu hoặc nốt vôi hóa của u bì, hoặc các ca u có cấu trúc đa vách phức tạp. Hiện tượng mất tín hiệu âm học phía sau khiến mặt nạ dự đoán của mô hình bị đứt quãng hoặc phân đoạn thiếu.")

    # Insert Image 5
    img5_path = "evaluation/baseline_visualizations/worst_matches.png"
    if os.path.exists(img5_path):
        p_img5 = doc.add_paragraph()
        p_img5.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img5.paragraph_format.first_line_indent = Cm(0)
        p_img5.paragraph_format.space_before = Pt(6)
        p_img5.paragraph_format.space_after = Pt(2)
        p_img5.add_run().add_picture(img5_path, width=Inches(4.8))
        add_figure_caption("Hình 5. Các trường hợp phân đoạn sai lệch do bóng cản âm hoặc cấu trúc u phức tạp")

    # -------------------------------------------------------------
    # SECTION V
    # -------------------------------------------------------------
    add_heading_1("V. KẾT QUẢ KIỂM THỬ TỰ ĐỘNG TOÀN DIỆN (AUTOMATED AUDIT SUITE)")
    add_p("Để bảo đảm tính toàn vẹn kỹ thuật, chống lỗi ngầm và xác nhận độ tin cậy của toàn bộ pipeline từ tiền xử lý đến mô hình, một bộ kiểm thử tự động gồm 16 kịch bản đã được xây dựng bằng thư viện PyTest (tại tests/test_milestone_1_audit.py):")

    add_table_title("Bảng 6. Danh mục 16 tiêu chí kiểm thử tự động toàn diện của hệ thống")
    t6_data = [
        ["STT", "Tên hàm kiểm thử", "Mục tiêu kiểm định kỹ thuật và nghiệp vụ", "Kết quả"],
        ["1", "test_dataset_discovery", "Xác nhận sự tồn tại và định dạng các tệp CSV phân chia dữ liệu", "Đạt (Pass)"],
        ["2", "test_image_mask_matching", "Kiểm tra việc khớp cặp 100% giữa tệp ảnh gốc và tệp mặt nạ tổn thương", "Đạt (Pass)"],
        ["3", "test_patient_level_split_leakage", "Kiểm tra độc lập tuyệt đối giữa các tập dữ liệu, không có bệnh nhân trùng lặp", "Đạt (Pass)"],
        ["4", "test_preprocessing_shape_consistency", "Xác nhận kích thước đầu ra cố định 512x512 sau phép biến đổi Letterbox", "Đạt (Pass)"],
        ["5", "test_mask_label_integrity", "Kiểm tra việc bảo toàn tập nhãn nhị phân {0, 1} của mặt nạ qua phép nội suy", "Đạt (Pass)"],
        ["6", "test_post_preprocessing_image_mask_pair", "Kiểm tra kiểu dữ liệu float32 và khoảng giá trị chuẩn hóa [0, 1]", "Đạt (Pass)"],
        ["7", "test_dataloader_batch", "Kiểm tra cấu trúc tensor đầu ra theo lô của PyTorch DataLoader", "Đạt (Pass)"],
        ["8", "test_unet_forward_pass", "Kiểm tra lan truyền tiến qua mô hình U-Net, không xuất hiện giá trị NaN", "Đạt (Pass)"],
        ["9", "test_loss_computation", "Kiểm tra việc tính toán hàm mất mát Combo Loss và đồ thị đạo hàm", "Đạt (Pass)"],
        ["10", "test_training_smoke_test", "Kiểm tra một bước cập nhật trọng số hoàn chỉnh của thuật toán tối ưu hóa", "Đạt (Pass)"],
        ["11", "test_prediction_shape", "Kiểm tra kích thước tensor dự đoán và ngưỡng phân lớp nhị phân 0,5", "Đạt (Pass)"],
        ["12", "test_predicted_mask_validity", "Kiểm tra tính đúng đắn của hàm ánh xạ ngược về kích thước ban đầu", "Đạt (Pass)"],
        ["13", "test_dice_calculation", "Kiểm tra tính toán chỉ số Dice trên dữ liệu kiểm soát giả lập", "Đạt (Pass)"],
        ["14", "test_iou_calculation", "Kiểm tra tính toán chỉ số IoU trên dữ liệu kiểm soát giả lập", "Đạt (Pass)"],
        ["15", "test_recall_calculation", "Kiểm tra tính toán độ nhạy Recall trên dữ liệu kiểm soát giả lập", "Đạt (Pass)"],
        ["16", "test_end_to_end_pipeline_flow", "Kiểm tra tích hợp toàn bộ chuỗi xử lý từ ảnh thô đến chỉ số đánh giá", "Đạt (Pass)"]
    ]
    table6 = doc.add_table(rows=len(t6_data), cols=4)
    table6.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table6)
    col_widths6 = [Cm(1.0), Cm(5.2), Cm(7.6), Cm(2.2)]
    
    for r_idx, row in enumerate(table6.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths6[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p_c = cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c_idx in [0, 3] or r_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
            p_c.paragraph_format.first_line_indent = Cm(0)
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(0)
            p_c.paragraph_format.line_spacing = 1.15
            run_c = p_c.add_run(t6_data[r_idx][c_idx])
            run_c.font.name = 'Times New Roman'
            run_c.font.size = Pt(10) if r_idx > 0 else Pt(10.5)
            run_c.font.bold = (r_idx == 0 or c_idx in [0, 3])
            if r_idx == 0:
                set_cell_shading(cell, "EDF2F7")
    
    add_table_source("Nguồn: Thực thi lệnh pytest tests/test_milestone_1_audit.py (Kết quả: 16 passed in 7.31s).")

    # -------------------------------------------------------------
    # SECTION VI
    # -------------------------------------------------------------
    add_heading_1("VI. NHẬN ĐỊNH KỸ THUẬT VÀ CÁC HẠN CHẾ CẦN GIẢI QUYẾT")
    add_p("Dựa trên các kết quả thực nghiệm đạt được ở Cột mốc 1, một số vấn đề kỹ thuật và thực tế nghiệp vụ được rút ra như sau:")
    add_p("1. Hiện tượng dự đoán lan rộng tại vùng biên mờ (Over-segmentation): Chỉ số Precision (58,6%) thấp hơn Recall (63,6%) cho thấy mô hình Standard U-Net có xu hướng phân đoạn thừa tại các ranh giới mô có độ phản hồi âm thấp tương tự dịch u. Do kiến trúc U-Net cơ sở sử dụng các phép tích chập cục bộ thuần túy mà chưa có cơ chế chọn lọc đặc trưng không gian, thông tin nhiễu từ các tầng mã hóa ban đầu truyền thẳng sang các tầng giải mã qua đường liên kết tắt.")
    add_p("2. Suy giảm tín hiệu do bóng cản âm (Acoustic Shadowing): Đối với các khối u bì buồng trứng chứa nhiều thành phần cản âm như nốt vôi hóa hoặc tóc, chùm sóng siêu âm bị triệt tiêu tạo thành dải bóng sẫm màu phía sau. Mô hình hiện tại gặp khó khăn trong việc suy luận hình học liên tục của khối u qua các dải bóng cản này.")
    add_p("3. Tính tất yếu của cơ chế tương tác Human-in-the-Loop: Kết quả định lượng trên tập kiểm thử khẳng định rằng mô hình học sâu đóng vai trò hỗ trợ gợi ý ban đầu chứ không thể thay thế việc xác nhận của bác sĩ chẩn đoán hình ảnh. Do đó, việc xây dựng công cụ tương tác trực quan (cho phép bác sĩ xem lớp phủ, dùng chổi vẽ hoặc tẩy xóa để tinh chỉnh mặt nạ trực tiếp trên trình duyệt) là mắt xích bắt buộc để bảo đảm độ chính xác và an toàn trước khi kết quả được lưu trữ hoặc đưa vào bệnh án.")

    # -------------------------------------------------------------
    # SECTION VII
    # -------------------------------------------------------------
    add_heading_1("VII. KẾ HOẠCH TRIỂN KHAI CỘT MỐC 2 (21/09/2026 - 05/10/2026)")
    add_p("Trong giai đoạn tiếp theo của đề tài, các công việc sẽ tập trung vào việc thử nghiệm mô hình Attention U-Net nhằm khắc phục hiện tượng phân đoạn thừa, đồng thời lập trình giao diện Web Prototype Human-in-the-Loop hoàn chỉnh:")

    add_table_title("Bảng 7. Kế hoạch công việc chi tiết trong Cột mốc 2")
    t7_data = [
        ["Thời gian", "Nhiệm vụ trọng tâm", "Nội dung kỹ thuật cụ thể", "Sản phẩm bàn giao dự kiến"],
        ["21/09 - 25/09", "Thử nghiệm kiến trúc Attention U-Net",
         "Tích hợp các cổng chú ý (Attention Gates) vào các đường liên kết tắt nhằm hạn chế nhiễu nền từ encoder, tập trung vào vùng tổn thương mục tiêu để nâng cao chỉ số Precision và Dice.",
         "Mã nguồn backend/models/attention_unet.py\nCheckpoint checkpoints/best_attention_unet.pth\nBảng thực nghiệm so sánh Ablation Study"],
        ["26/09 - 28/09", "Đánh giá so sánh và phân tích lỗi",
         "Đánh giá đối chứng Standard U-Net và Attention U-Net trên cùng tập kiểm thử 46 ca; phân tích lỗi chi tiết trên các nhóm ca bóng cản âm và ca u vách ngăn; lựa chọn mô hình chính thức.",
         "Báo cáo đối chứng định lượng (model_comparison.csv)\nBộ hình ảnh phân tích lỗi chuyên sâu"],
        ["29/09 - 30/09", "Xây dựng dịch vụ suy luận (Inference API)",
         "Phát triển RESTful API bằng FastAPI; tối ưu hóa pipeline tiền xử lý ảnh và nạp mô hình; bảo đảm thời gian đáp ứng dưới 500 mili-giây mỗi ảnh trên phần cứng GPU cá nhân.",
         "Mã nguồn backend/api/segmentation.py\nĐiểm cuối /api/v1/segmentation/predict hoạt động ổn định"],
        ["01/10 - 03/10", "Xây dựng giao diện Web Prototype HITL",
         "Lập trình giao diện người dùng bằng React.js và HTML5 Canvas; cung cấp tính năng tải ảnh, hiển thị lớp phủ mặt nạ gợi ý, công cụ chổi vẽ (Brush) và tẩy xóa (Eraser) để bác sĩ chỉnh sửa và xác nhận kết quả.",
         "Giao diện Web Prototype tương tác trực tiếp trên trình duyệt"],
        ["04/10 - 05/10", "Tổng hợp kết quả và soạn thảo báo cáo",
         "Tổng hợp dữ liệu thực nghiệm đưa vào bản thảo Chương 3 khóa luận tốt nghiệp; hoàn thiện Báo cáo tiến độ Cột mốc 2 nộp Thầy.",
         "Bản thảo Chương 3 KLTN\nBáo cáo tiến độ Cột mốc 2"]
    ]
    table7 = doc.add_table(rows=len(t7_data), cols=4)
    table7.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table7)
    col_widths7 = [Cm(2.6), Cm(4.0), Cm(5.6), Cm(3.8)]
    
    for r_idx, row in enumerate(table7.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths7[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p_c = cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c_idx == 0 or r_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
            p_c.paragraph_format.first_line_indent = Cm(0)
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(0)
            p_c.paragraph_format.line_spacing = 1.15
            run_c = p_c.add_run(t7_data[r_idx][c_idx])
            run_c.font.name = 'Times New Roman'
            run_c.font.size = Pt(10) if r_idx > 0 else Pt(10.5)
            run_c.font.bold = (r_idx == 0 or c_idx == 0)
            if r_idx == 0:
                set_cell_shading(cell, "EDF2F7")
    
    add_table_source("Nguồn: Kế hoạch triển khai đề cương khóa luận tốt nghiệp MIS 65A.")

    add_p("Lưu ý về phạm vi triển khai: Bám sát chỉ đạo chuyên môn của Thầy, các hướng nghiên cứu mở rộng như tích hợp mô hình ngôn ngữ lớn (LLM/VLM), tự động sinh văn bản mô tả bệnh án hoặc phát triển bảng điều khiển quản trị nâng cao sẽ chỉ được xem xét nếu phần phân đoạn và giao diện tương tác Human-in-the-Loop cốt lõi đã hoàn thành tốt và còn đủ thời gian.")

    # Closing
    add_p("Em xin kính báo cáo Thầy tình hình thực hiện Cột mốc 1 của khóa luận. Em rất mong tiếp tục nhận được những lời khuyên và chỉ dẫn chuyên môn của Thầy để thực hiện tốt các nội dung nghiên cứu trong Cột mốc 2.")
    add_p("Em xin chân thành cảm ơn Thầy.", bold=False)

    # Signature
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.paragraph_format.first_line_indent = Cm(0)
    p_sig.paragraph_format.space_before = Pt(12)
    p_sig.paragraph_format.space_after = Pt(2)
    r_sig_lbl = p_sig.add_run("Sinh viên thực hiện\n\n\n\n")
    r_sig_lbl.font.name = 'Times New Roman'
    r_sig_lbl.font.size = Pt(12)
    r_sig_lbl.font.bold = True
    
    r_sig_name = p_sig.add_run("Nguyễn Hữu Dũng\n")
    r_sig_name.font.name = 'Times New Roman'
    r_sig_name.font.size = Pt(13)
    r_sig_name.font.bold = True

    r_sig_id = p_sig.add_run("Mã sinh viên: 11235559, Lớp: HTTTQL 65A")
    r_sig_id.font.name = 'Times New Roman'
    r_sig_id.font.size = Pt(11)
    r_sig_id.font.italic = True

    # Save to both target paths
    path1 = "docs/reports/Bao_Cao_Tien_Do_Moc_1_NguyenHuuDung.docx"
    path2 = "docs/reports/Bao_Cao_Tien_Do_Moc_1_NguyenHuuDung_Final.docx"
    
    doc.save(path1)
    print(f"Saved: {path1}")
    doc.save(path2)
    print(f"Saved: {path2}")

if __name__ == "__main__":
    create_report()
